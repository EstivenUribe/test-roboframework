"""
ExcelReader.py - Librería personalizada para leer datos de prueba desde Excel.
Usada en todos los suites de Robot Framework para el proyecto Biblioteca Pro.
"""
from openpyxl import load_workbook


class ExcelReader:
    """Lee datos de prueba desde hojas de Excel (.xlsx)."""

    ROBOT_LIBRARY_SCOPE = "GLOBAL"

    def __init__(self):
        self._workbook = None
        self._path = None

    def cargar_excel(self, ruta):
        """Carga el archivo Excel en memoria.

        Argumento:
        - ruta: ruta absoluta al archivo .xlsx
        """
        self._workbook = load_workbook(ruta, data_only=True)
        self._path = ruta
        return f"Excel cargado: {ruta}"

    def obtener_datos_hoja(self, nombre_hoja):
        """Devuelve una lista de diccionarios con los datos de la hoja.

        La primera fila se usa como encabezados (claves del diccionario).
        Las filas completamente vacías se omiten.
        """
        if self._workbook is None:
            raise RuntimeError("Primero llama a 'Cargar Excel'.")
        ws = self._workbook[nombre_hoja]
        encabezados = [str(c.value).strip() if c.value else "" for c in ws[1]]
        datos = []
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in fila):
                continue
            registro = {encabezados[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(fila)}
            datos.append(registro)
        return datos

    def obtener_fila_por_caso(self, nombre_hoja, id_caso):
        """Devuelve el diccionario de la fila cuyo campo 'id_caso' coincide."""
        filas = self.obtener_datos_hoja(nombre_hoja)
        for f in filas:
            if f.get("id_caso", "").upper() == str(id_caso).upper():
                return f
        raise ValueError(f"Caso '{id_caso}' no encontrado en hoja '{nombre_hoja}'.")

    def obtener_nombres_hojas(self):
        """Devuelve la lista de nombres de hojas del workbook."""
        if self._workbook is None:
            raise RuntimeError("Primero llama a 'Cargar Excel'.")
        return self._workbook.sheetnames
